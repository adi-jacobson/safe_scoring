from __future__ import annotations

import math
import shutil
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = Path(__file__).with_name("scoring_spreadsheet_2026.xlsx")
OUTPUT = Path(__file__).with_name("scoring_fairness_analysis.xlsx")
REPORT = Path(__file__).with_name("scoring_fairness_report.md")
WEIGHT_COLUMNS = {"judge": 5, "peer": 6, "attack": 7, "resistance": 8}
TEAM_START_ROW = 6
TEAM_CODE_COLUMN = 3
TEAM_NAME_COLUMN = 2
CRACK_HEADER_ROW = 4
CRACK_COLUMN_RANGE = range(44, 74)
PEER_HEADER_ROW = 4
PEER_COLUMN_RANGE = range(14, 44)
COMPONENT_COLUMNS = {"judge": 74, "peer": 75, "attack": 76, "resistance": 77, "total": 78}
VALID_OUTCOMES = {0, 0.5, 1}
INVALID_OUTCOME = 0.5
CONFIDENCE_Z = 1.959963984540054


def logistic(values: np.ndarray) -> np.ndarray:
    """Apply a numerically stable logistic transform."""
    values = np.clip(values, -30, 30)
    return 1.0 / (1.0 + np.exp(-values))


def required_float(value: object, label: str) -> float:
    """Return a numeric workbook value or fail with a useful message."""
    if value is None or isinstance(value, bool):
        raise ValueError(f"Missing numeric value for {label}. Recalculate and save the workbook in Excel.")
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Expected a numeric value for {label}, got {value!r}.") from error
    if not math.isfinite(result):
        raise ValueError(f"Expected a finite value for {label}, got {value!r}.")
    return result


def fit_rasch(
    design: np.ndarray,
    outcomes: np.ndarray,
    regularisation: float,
) -> tuple[np.ndarray, np.ndarray]:
    """Fit a ridge-regularised Rasch logistic model by Newton iteration."""
    parameters = np.zeros(design.shape[1])
    penalty = np.full(design.shape[1], regularisation)
    penalty[0] = 0.0

    for _ in range(100):
        probabilities = logistic(design @ parameters)
        variances = np.maximum(probabilities * (1 - probabilities), 1e-8)
        gradient = design.T @ (probabilities - outcomes) + penalty * parameters
        hessian = (design.T * variances) @ design + np.diag(penalty)
        step = np.linalg.solve(hessian, gradient)
        parameters -= step
        if np.max(np.abs(step)) < 1e-9:
            break
    else:
        raise RuntimeError("Rasch model did not converge within 100 Newton iterations.")

    covariance = np.linalg.solve(hessian, np.eye(hessian.shape[0]))
    return parameters, covariance


def make_design(
    observations: list[dict[str, object]],
    team_count: int,
) -> tuple[np.ndarray, np.ndarray]:
    """Build the attacker-minus-defender design matrix."""
    design = np.zeros((len(observations), 1 + 2 * team_count))
    outcomes = np.empty(len(observations))
    for row_index, observation in enumerate(observations):
        attacker = int(observation["attacker_index"])
        defender = int(observation["defender_index"])
        design[row_index, 0] = 1
        design[row_index, 1 + attacker] = 1
        design[row_index, 1 + team_count + defender] = -1
        outcomes[row_index] = float(observation["outcome"])
    return design, outcomes


def choose_regularisation(
    design: np.ndarray,
    outcomes: np.ndarray,
) -> tuple[float, dict[float, float]]:
    """Select ridge strength with deterministic five-fold log loss."""
    candidates = (0.1, 0.25, 0.5, 1.0, 2.0, 4.0, 8.0)
    rng = np.random.default_rng(2026)
    folds = np.arange(len(outcomes)) % 5
    rng.shuffle(folds)
    losses: dict[float, float] = {}

    for candidate in candidates:
        fold_losses = []
        for fold in range(5):
            train = folds != fold
            test = ~train
            parameters, _ = fit_rasch(design[train], outcomes[train], candidate)
            probabilities = np.clip(logistic(design[test] @ parameters), 1e-9, 1 - 1e-9)
            loss = -np.mean(
                outcomes[test] * np.log(probabilities)
                + (1 - outcomes[test]) * np.log(1 - probabilities)
            )
            fold_losses.append(float(loss))
        losses[candidate] = float(np.mean(fold_losses))

    best = min(losses, key=losses.get)
    return best, losses


def fit_peer_model(
    design: np.ndarray,
    scores: np.ndarray,
    regularisation: float,
) -> tuple[np.ndarray, np.ndarray, float]:
    """Fit a ridge additive model for recipient quality and rater generosity."""
    penalty = np.full(design.shape[1], regularisation)
    penalty[0] = 0.0
    precision = design.T @ design + np.diag(penalty)
    precision_inverse = np.linalg.solve(precision, np.eye(precision.shape[0]))
    parameters = precision_inverse @ design.T @ scores
    residuals = scores - design @ parameters
    effective_parameters = float(np.trace(precision_inverse @ design.T @ design))
    residual_degrees_of_freedom = max(len(scores) - effective_parameters, 1)
    residual_variance = float(np.sum(residuals**2) / residual_degrees_of_freedom)
    covariance = residual_variance * precision_inverse
    return parameters, covariance, residual_variance


def choose_peer_regularisation(
    design: np.ndarray,
    scores: np.ndarray,
) -> tuple[float, dict[float, float]]:
    """Select peer-model ridge strength with deterministic five-fold RMSE."""
    candidates = (0.1, 0.25, 0.5, 1.0, 2.0, 4.0, 8.0, 16.0)
    rng = np.random.default_rng(2026)
    folds = np.arange(len(scores)) % 5
    rng.shuffle(folds)
    losses: dict[float, float] = {}
    for candidate in candidates:
        squared_errors = []
        for fold in range(5):
            train = folds != fold
            test = ~train
            parameters, _, _ = fit_peer_model(design[train], scores[train], candidate)
            squared_errors.extend((scores[test] - design[test] @ parameters) ** 2)
        losses[candidate] = float(np.sqrt(np.mean(squared_errors)))
    return min(losses, key=losses.get), losses


def peer_mean_baseline_rmse(scores: np.ndarray) -> float:
    """Return deterministic five-fold RMSE for a training-mean predictor."""
    rng = np.random.default_rng(2026)
    folds = np.arange(len(scores)) % 5
    rng.shuffle(folds)
    squared_errors = []
    for fold in range(5):
        train = folds != fold
        test = ~train
        squared_errors.extend((scores[test] - np.mean(scores[train])) ** 2)
    return float(np.sqrt(np.mean(squared_errors)))


def make_peer_design(
    observations: list[dict[str, object]],
    team_count: int,
) -> tuple[np.ndarray, np.ndarray]:
    """Build an intercept-plus-recipient-plus-rater additive design matrix."""
    design = np.zeros((len(observations), 1 + 2 * team_count))
    scores = np.empty(len(observations))
    for row_index, observation in enumerate(observations):
        recipient = int(observation["recipient_index"])
        rater = int(observation["rater_index"])
        design[row_index, 0] = 1
        design[row_index, 1 + recipient] = 1
        design[row_index, 1 + team_count + rater] = 1
        scores[row_index] = float(observation["score"])
    return design, scores


def rank_descending(values: list[float]) -> list[int]:
    """Assign dense descending ranks."""
    ordered = sorted(set(values), reverse=True)
    positions = {value: index + 1 for index, value in enumerate(ordered)}
    return [positions[value] for value in values]


def load_workbook_copy(source: Path = SOURCE) -> openpyxl.Workbook:
    """Open cached workbook values, copying first when Excel locks the file."""
    try:
        return openpyxl.load_workbook(source, data_only=True, read_only=True)
    except PermissionError:
        temporary_path = Path(tempfile.gettempdir()) / source.name
        shutil.copy2(source, temporary_path)
        return openpyxl.load_workbook(temporary_path, data_only=True, read_only=True)


def analyse(source: Path = SOURCE) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, object]]:
    """Validate the workbook and produce adjusted scores and diagnostics."""
    workbook = load_workbook_copy(source)
    missing_sheets = {"Tracker", "Global Values"} - set(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(f"Workbook is missing required sheets: {sorted(missing_sheets)}")
    tracker = workbook["Tracker"]
    globals_sheet = workbook["Global Values"]
    weights = {
        name: required_float(
            globals_sheet.cell(row=WEIGHT_COLUMNS[name], column=2).value,
            f"Global Values!B{WEIGHT_COLUMNS[name]} ({name} weight)",
        )
        for name in WEIGHT_COLUMNS
    }
    if any(weight < 0 for weight in weights.values()):
        raise ValueError(f"Scoring weights must be non-negative: {weights}")
    if not math.isclose(sum(weights.values()), 1.0, abs_tol=1e-9):
        raise ValueError(f"Scoring weights must sum to 1.0, got {sum(weights.values()):.6f}.")

    team_rows = [
        row
        for row in range(TEAM_START_ROW, tracker.max_row + 1)
        if tracker.cell(row=row, column=TEAM_CODE_COLUMN).value
    ]
    teams = [str(tracker.cell(row=row, column=TEAM_NAME_COLUMN).value) for row in team_rows]
    codes = [str(tracker.cell(row=row, column=TEAM_CODE_COLUMN).value) for row in team_rows]
    if len(teams) < 2:
        raise ValueError("At least two scored teams are required.")
    if len(set(teams)) != len(teams) or len(set(codes)) != len(codes):
        raise ValueError("Team names and codes must be unique.")
    team_lookup = {team: index for index, team in enumerate(teams)}

    max_peer_score = required_float(globals_sheet.cell(row=3, column=2).value, "Global Values!B3 (max peer score)")
    if max_peer_score <= 0:
        raise ValueError("Maximum peer score must be positive.")

    peer_columns = []
    for column in PEER_COLUMN_RANGE:
        header = tracker.cell(row=PEER_HEADER_ROW, column=column).value
        if header in team_lookup:
            peer_columns.append((column, team_lookup[str(header)]))
    if len(peer_columns) != len(teams):
        raise ValueError(
            f"Expected one peer-rating column per team; found {len(peer_columns)} columns for {len(teams)} teams."
        )

    peer_observations: list[dict[str, object]] = []
    for recipient_index, row in enumerate(team_rows):
        for column, rater_index in peer_columns:
            value = tracker.cell(row=row, column=column).value
            if value is None:
                continue
            score = required_float(value, f"peer score in Tracker!{tracker.cell(row=row, column=column).coordinate}")
            if not 0 <= score <= max_peer_score:
                coordinate = tracker.cell(row=row, column=column).coordinate
                raise ValueError(f"Peer score in Tracker!{coordinate} must be between 0 and {max_peer_score:g}.")
            if rater_index == recipient_index:
                coordinate = tracker.cell(row=row, column=column).coordinate
                raise ValueError(f"Self-rating recorded in Tracker!{coordinate} for {teams[rater_index]}.")
            peer_observations.append(
                {
                    "recipient_index": recipient_index,
                    "rater_index": rater_index,
                    "recipient": teams[recipient_index],
                    "rater": teams[rater_index],
                    "score": score / max_peer_score,
                }
            )
    for team_index, code in enumerate(codes):
        ratings_given = sum(a["rater_index"] == team_index for a in peer_observations)
        ratings_received = sum(a["recipient_index"] == team_index for a in peer_observations)
        if ratings_given == 0 or ratings_received == 0:
            raise ValueError(f"Team {code} needs at least one peer rating given and received.")

    peer_design, peer_outcomes = make_peer_design(peer_observations, len(teams))
    peer_regularisation, peer_cv_losses = choose_peer_regularisation(peer_design, peer_outcomes)
    peer_parameters, peer_covariance, peer_residual_variance = fit_peer_model(
        peer_design, peer_outcomes, peer_regularisation
    )
    peer_intercept = peer_parameters[0]
    peer_quality_effects = peer_parameters[1 : 1 + len(teams)]
    peer_rater_effects = peer_parameters[1 + len(teams) :]
    adjusted_peer_scores = np.clip(peer_intercept + peer_quality_effects, 0, 1)
    peer_vectors = []
    peer_intervals = []
    for team_index in range(len(teams)):
        peer_vector = np.zeros(len(peer_parameters))
        peer_vector[0] = 1
        peer_vector[1 + team_index] = 1
        peer_vectors.append(peer_vector)
        peer_se = math.sqrt(float(peer_vector @ peer_covariance @ peer_vector))
        peer_intervals.append(
            (
                float(max(0, adjusted_peer_scores[team_index] - CONFIDENCE_Z * peer_se)),
                float(min(1, adjusted_peer_scores[team_index] + CONFIDENCE_Z * peer_se)),
            )
        )

    crack_columns = []
    for column in CRACK_COLUMN_RANGE:
        header = tracker.cell(row=CRACK_HEADER_ROW, column=column).value
        if header in team_lookup:
            crack_columns.append((column, team_lookup[str(header)]))
    if len(crack_columns) != len(teams):
        raise ValueError(
            f"Expected one cracking column per team; found {len(crack_columns)} columns for {len(teams)} teams."
        )

    all_attempts: list[dict[str, object]] = []
    valid_observations: list[dict[str, object]] = []
    for defender_index, row in enumerate(team_rows):
        for column, attacker_index in crack_columns:
            value = tracker.cell(row=row, column=column).value
            if value is None:
                continue
            if value not in VALID_OUTCOMES:
                coordinate = tracker.cell(row=row, column=column).coordinate
                raise ValueError(f"Unexpected cracking result {value!r} in Tracker!{coordinate}; use 0, 0.5, 1, or blank.")
            if attacker_index == defender_index:
                coordinate = tracker.cell(row=row, column=column).coordinate
                raise ValueError(f"Self-match recorded in Tracker!{coordinate} for {teams[attacker_index]}.")
            attempt = {
                "attacker_index": attacker_index,
                "defender_index": defender_index,
                "attacker": teams[attacker_index],
                "defender": teams[defender_index],
                "outcome": float(value),
                "status": "Invalid/unreliable" if value == INVALID_OUTCOME else "Valid",
            }
            all_attempts.append(attempt)
            if value != INVALID_OUTCOME:
                valid_observations.append(attempt)

    if not valid_observations:
        raise ValueError("No valid 0/1 cracking outcomes were found.")
    for team_index, code in enumerate(codes):
        valid_attacks = sum(a["attacker_index"] == team_index for a in valid_observations)
        valid_defences = sum(a["defender_index"] == team_index for a in valid_observations)
        if valid_attacks == 0 or valid_defences == 0:
            raise ValueError(f"Team {code} needs at least one valid attack and defence result.")

    design, outcomes = make_design(valid_observations, len(teams))
    regularisation, cv_losses = choose_regularisation(design, outcomes)
    parameters, covariance = fit_rasch(design, outcomes, regularisation)
    intercept = parameters[0]
    attack_effects = parameters[1 : 1 + len(teams)]
    defence_effects = parameters[1 + len(teams) :]

    attack_scores = logistic(intercept + attack_effects)
    resistance_scores = logistic(defence_effects - intercept)
    attack_intervals = []
    resistance_intervals = []
    total_margins = []
    total_rasch_gradients = []
    total_peer_gradients = []
    for team_index in range(len(teams)):
        attack_vector = np.zeros(len(parameters))
        attack_vector[0] = 1
        attack_vector[1 + team_index] = 1
        attack_eta = intercept + attack_effects[team_index]
        attack_se = math.sqrt(float(attack_vector @ covariance @ attack_vector))
        attack_intervals.append(
            tuple(float(logistic(np.array([attack_eta + sign * CONFIDENCE_Z * attack_se]))[0]) for sign in (-1, 1))
        )

        resistance_vector = np.zeros(len(parameters))
        resistance_vector[0] = -1
        resistance_vector[1 + len(teams) + team_index] = 1
        resistance_eta = defence_effects[team_index] - intercept
        resistance_se = math.sqrt(float(resistance_vector @ covariance @ resistance_vector))
        resistance_intervals.append(
            tuple(
                float(logistic(np.array([resistance_eta + sign * CONFIDENCE_Z * resistance_se]))[0])
                for sign in (-1, 1)
            )
        )

        total_gradient = 100 * (
            weights["attack"]
            * attack_scores[team_index]
            * (1 - attack_scores[team_index])
            * attack_vector
            + weights["resistance"]
            * resistance_scores[team_index]
            * (1 - resistance_scores[team_index])
            * resistance_vector
        )
        peer_gradient = 100 * weights["peer"] * peer_vectors[team_index]
        total_variance = float(
            total_gradient @ covariance @ total_gradient
            + peer_gradient @ peer_covariance @ peer_gradient
        )
        total_se = math.sqrt(total_variance)
        total_margins.append(CONFIDENCE_Z * total_se)
        total_rasch_gradients.append(total_gradient)
        total_peer_gradients.append(peer_gradient)

    rows: list[dict[str, object]] = []
    for team_index, (row, team, code) in enumerate(zip(team_rows, teams, codes)):
        attack_attempts = [a for a in all_attempts if a["attacker_index"] == team_index]
        defence_attempts = [a for a in all_attempts if a["defender_index"] == team_index]
        valid_attacks = [a for a in attack_attempts if a["status"] == "Valid"]
        valid_defences = [a for a in defence_attempts if a["status"] == "Valid"]
        faced_difficulties = [
            float(logistic(np.array([intercept - defence_effects[int(a["defender_index"])] ]))[0])
            for a in attack_attempts
        ]
        faced_attackers = [
            float(logistic(np.array([intercept + attack_effects[int(a["attacker_index"])] ]))[0])
            for a in defence_attempts
        ]
        judge = required_float(tracker.cell(row=row, column=COMPONENT_COLUMNS["judge"]).value, f"{code} judge score")
        peer = required_float(tracker.cell(row=row, column=COMPONENT_COLUMNS["peer"]).value, f"{code} peer score")
        current_attack = required_float(tracker.cell(row=row, column=COMPONENT_COLUMNS["attack"]).value, f"{code} crack score")
        current_resistance = required_float(tracker.cell(row=row, column=COMPONENT_COLUMNS["resistance"]).value, f"{code} resistance score")
        current_total = required_float(tracker.cell(row=row, column=COMPONENT_COLUMNS["total"]).value, f"{code} total score")
        adjusted_total = 100 * (
            judge * weights["judge"]
            + adjusted_peer_scores[team_index] * weights["peer"]
            + attack_scores[team_index] * weights["attack"]
            + resistance_scores[team_index] * weights["resistance"]
        )
        rows.append(
            {
                "code": code,
                "team": team,
                "judge": judge,
                "peer": peer,
                "adjusted_peer": float(adjusted_peer_scores[team_index]),
                "peer_low": peer_intervals[team_index][0],
                "peer_high": peer_intervals[team_index][1],
                "peer_rater_effect": float(peer_rater_effects[team_index]),
                "peer_ratings_given": sum(a["rater_index"] == team_index for a in peer_observations),
                "peer_ratings_received": sum(a["recipient_index"] == team_index for a in peer_observations),
                "current_attack": current_attack,
                "valid_attack": float(np.mean([a["outcome"] for a in valid_attacks])),
                "adjusted_attack": float(attack_scores[team_index]),
                "attack_low": attack_intervals[team_index][0],
                "attack_high": attack_intervals[team_index][1],
                "current_resistance": current_resistance,
                "valid_resistance": float(1 - np.mean([a["outcome"] for a in valid_defences])),
                "adjusted_resistance": float(resistance_scores[team_index]),
                "resistance_low": resistance_intervals[team_index][0],
                "resistance_high": resistance_intervals[team_index][1],
                "valid_attacks": len(valid_attacks),
                "invalid_attacks": len(attack_attempts) - len(valid_attacks),
                "valid_defences": len(valid_defences),
                "invalid_defences": len(defence_attempts) - len(valid_defences),
                "schedule_easiness": float(np.mean(faced_difficulties)),
                "attacker_strength_faced": float(np.mean(faced_attackers)),
                "current_total": current_total,
                "adjusted_total": float(adjusted_total),
                "adjusted_total_low": float(max(0, adjusted_total - total_margins[team_index])),
                "adjusted_total_high": float(min(100, adjusted_total + total_margins[team_index])),
            }
        )

    current_ranks = rank_descending([float(row["current_total"]) for row in rows])
    adjusted_ranks = rank_descending([float(row["adjusted_total"]) for row in rows])
    for row, current_rank, adjusted_rank in zip(rows, current_ranks, adjusted_ranks):
        row["current_rank"] = current_rank
        row["adjusted_rank"] = adjusted_rank
        row["rank_change"] = current_rank - adjusted_rank

    leader_index = adjusted_ranks.index(1)
    leader_score = float(rows[leader_index]["adjusted_total"])
    for team_index, row in enumerate(rows):
        difference = float(row["adjusted_total"]) - leader_score
        rasch_difference_gradient = total_rasch_gradients[team_index] - total_rasch_gradients[leader_index]
        peer_difference_gradient = total_peer_gradients[team_index] - total_peer_gradients[leader_index]
        difference_variance = float(
            rasch_difference_gradient @ covariance @ rasch_difference_gradient
            + peer_difference_gradient @ peer_covariance @ peer_difference_gradient
        )
        difference_se = math.sqrt(difference_variance)
        difference_margin = CONFIDENCE_Z * difference_se
        row["difference_from_leader"] = difference
        row["difference_from_leader_low"] = difference - difference_margin
        row["difference_from_leader_high"] = difference + difference_margin
        row["winner_review"] = team_index == leader_index or difference + difference_margin >= 0

    diagnostics = {
        "weights": weights,
        "team_count": len(teams),
        "attempt_count": len(all_attempts),
        "valid_count": len(valid_observations),
        "invalid_count": len(all_attempts) - len(valid_observations),
        "regularisation": regularisation,
        "cv_losses": cv_losses,
        "overall_valid_crack_rate": float(np.mean(outcomes)),
        "baseline_log_loss": float(
            -np.mean(
                outcomes * np.log(np.mean(outcomes))
                + (1 - outcomes) * np.log(1 - np.mean(outcomes))
            )
        ),
        "peer_regularisation": peer_regularisation,
        "peer_cv_losses": peer_cv_losses,
        "peer_cv_rmse": peer_cv_losses[peer_regularisation],
        "peer_baseline_rmse": peer_mean_baseline_rmse(peer_outcomes),
        "peer_residual_sd": math.sqrt(peer_residual_variance),
        "interval_method": "Laplace approximation conditional on the selected regularisation",
    }
    workbook.close()
    return rows, all_attempts, diagnostics


def add_table_style(sheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 48)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def write_outputs(
    rows: list[dict[str, object]],
    attempts: list[dict[str, object]],
    diagnostics: dict[str, object],
    output_path: Path = OUTPUT,
    report_path: Path = REPORT,
) -> None:
    workbook = Workbook()
    decision = workbook.active
    decision.title = "Decision table"
    decision_headers = [
        "Fair rank", "Code", "Team / safe", "Raw total", "Judges /45", "Peers /20",
        "Cracking /25", "Resistance /10", "Fair score", "Fair score range", "Status",
    ]
    decision.append(decision_headers)
    ordered_rows = sorted(rows, key=lambda row: int(row["adjusted_rank"]))
    for row in ordered_rows:
        if row["adjusted_rank"] == 1:
            status = "Provisional leader - moderate"
        elif row["winner_review"]:
            status = "Winner review"
        else:
            status = "Outside winner-review group"
        decision.append([
            row["adjusted_rank"], row["code"], row["team"], row["current_total"],
            100 * row["judge"] * diagnostics["weights"]["judge"],
            100 * row["adjusted_peer"] * diagnostics["weights"]["peer"],
            100 * row["adjusted_attack"] * diagnostics["weights"]["attack"],
            100 * row["adjusted_resistance"] * diagnostics["weights"]["resistance"],
            row["adjusted_total"],
            f"{row['adjusted_total_low']:.1f} - {row['adjusted_total_high']:.1f}",
            status,
        ])
    for row in decision.iter_rows(min_row=2, min_col=4, max_col=9):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    add_table_style(decision)

    summary = workbook.create_sheet("Audit detail")
    headers = [
        "Fair rank", "Current rank", "Change", "Code", "Team / safe", "Judge", "Raw peer",
        "Adjusted peer", "Peer 95% low", "Peer 95% high", "Rater generosity",
        "Current crack", "Valid-only crack", "Adjusted crack", "Crack 95% low", "Crack 95% high",
        "Current resistance", "Valid-only resistance", "Adjusted resistance", "Resistance 95% low",
        "Resistance 95% high", "Valid attacks", "Invalid attacks", "Valid defences", "Invalid defences",
        "Schedule easiness", "Attacker strength faced", "Current total", "Fair total",
        "Fair total 95% low", "Fair total 95% high", "Winner review?",
    ]
    summary.append(headers)
    for row in ordered_rows:
        summary.append([
            row["adjusted_rank"], row["current_rank"], row["rank_change"], row["code"], row["team"],
            row["judge"], row["peer"], row["adjusted_peer"], row["peer_low"], row["peer_high"],
            row["peer_rater_effect"], row["current_attack"], row["valid_attack"], row["adjusted_attack"],
            row["attack_low"], row["attack_high"], row["current_resistance"], row["valid_resistance"],
            row["adjusted_resistance"], row["resistance_low"], row["resistance_high"], row["valid_attacks"],
            row["invalid_attacks"], row["valid_defences"], row["invalid_defences"], row["schedule_easiness"],
            row["attacker_strength_faced"], row["current_total"], row["adjusted_total"],
            row["adjusted_total_low"], row["adjusted_total_high"],
            "Yes" if row["winner_review"] else "No",
        ])
    for row in summary.iter_rows(min_row=2, min_col=6, max_col=32):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000"
    add_table_style(summary)

    chart = BarChart()
    chart.title = "Fair total by team"
    chart.y_axis.title = "Points"
    chart.height = 9
    chart.width = 18
    chart.add_data(Reference(decision, min_col=9, min_row=1, max_row=decision.max_row), titles_from_data=True)
    chart.set_categories(Reference(decision, min_col=2, min_row=2, max_row=decision.max_row))
    decision.add_chart(chart, "M2")

    attempts_sheet = workbook.create_sheet("Attempt audit")
    attempts_sheet.append(["Attacker", "Defending safe", "Recorded value", "Treatment"])
    for attempt in attempts:
        attempts_sheet.append([attempt["attacker"], attempt["defender"], attempt["outcome"], attempt["status"]])
    add_table_style(attempts_sheet)

    peer_sheet = workbook.create_sheet("Peer rater audit")
    peer_sheet.append(["Code", "Team", "Rater generosity (points)", "Ratings given", "Ratings received"])
    for row in sorted(rows, key=lambda item: str(item["code"])):
        peer_sheet.append([
            row["code"], row["team"], 100 * float(row["peer_rater_effect"]),
            row["peer_ratings_given"], row["peer_ratings_received"],
        ])
    for cell in peer_sheet["C"][1:]:
        cell.number_format = "+0.0;-0.0;0.0"
    add_table_style(peer_sheet)

    method = workbook.create_sheet("Method and caveats")
    method_rows = [
        ("Purpose", "Rank teams while accounting for variation in the difficulty of safes attacked and strength of attackers faced."),
        ("Invalid records", "A recorded 0.5 is excluded from model fitting. It is not evidence of either success or failure."),
        ("Model", "Regularised Rasch logistic model: logit(P(crack)) = overall rate + attacker ability - safe difficulty."),
        ("Peer model", "Regularised additive model: peer score = overall level + safe quality + rater generosity + noise."),
        ("Adjusted peer", "Estimated peer score after accounting for whether each scoring team tends to be harsh or generous."),
        ("Adjusted crack", "Estimated chance that the team cracks an average-difficulty safe."),
        ("Adjusted resistance", "Estimated chance that the safe resists an average-strength attacker."),
        ("Fair total", "45% judge + 20% adjusted peer + 25% adjusted crack + 10% adjusted resistance."),
        ("Regularisation", f"Selected by deterministic 5-fold cross-validation: crack model {diagnostics['regularisation']}; peer model {diagnostics['peer_regularisation']}."),
        ("Uncertainty", "Approximate 95% model intervals are shown. Overlapping intervals mean the crack evidence does not strongly separate teams."),
        ("Winner review", "Yes means the modelled score difference from the provisional leader includes zero at the 95% level. These teams should be moderated together before declaring a winner."),
        ("Scope of intervals", "Fair-total intervals include conditional model uncertainty in adjusted peer, cracking, and resistance scores. Judge scores are treated as fixed."),
        ("Important limitation", "The model reduces schedule luck; it cannot recover information from invalid attempts or remove all uncertainty from only a few valid attempts."),
        ("Operational rule", "Do not use adjusted crack metrics as a tie-breaker when intervals overlap substantially. Use a declared practical tie-break or judging moderation."),
    ]
    for label, detail in method_rows:
        method.append([label, detail])
    method.column_dimensions["A"].width = 24
    method.column_dimensions["B"].width = 110
    method.freeze_panes = "A2"
    for cell in method[1]:
        cell.font = Font(bold=True)
    for row in method.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(output_path)

    biggest_movers = sorted(rows, key=lambda row: abs(int(row["rank_change"])), reverse=True)[:5]
    schedule_sorted = sorted(rows, key=lambda row: float(row["schedule_easiness"]))
    report_lines = [
        "# Safe-cracking scoring fairness analysis",
        "",
        "## Recommendation",
        "",
        "Treat `0.5` as missing evidence, not half a crack. Adjust peer scores for rater harshness, and adjust cracking and resistance for opponent difficulty, while retaining the published 45% / 20% / 25% / 10% weights.",
        "",
        "## Data checks",
        "",
        f"- {diagnostics['team_count']} teams, {diagnostics['attempt_count']} recorded matchups.",
        f"- {diagnostics['valid_count']} valid binary outcomes and {diagnostics['invalid_count']} invalid/unreliable `0.5` records.",
        f"- Overall valid crack rate: {diagnostics['overall_valid_crack_rate']:.1%}.",
        f"- Cross-validated regularisation strength: {diagnostics['regularisation']}.",
        f"- Cross-validated log loss: {diagnostics['cv_losses'][diagnostics['regularisation']]:.3f}; intercept-only baseline: {diagnostics['baseline_log_loss']:.3f}.",
        f"- Peer regularisation strength: {diagnostics['peer_regularisation']}; cross-validated RMSE: {diagnostics['peer_cv_rmse']:.3f}; mean-only baseline: {diagnostics['peer_baseline_rmse']:.3f}.",
        "",
        "## Adjusted ranking",
        "",
        "| Rank | Code | Team / safe | Raw total | Judges /45 | Peers /20 | Crack /25 | Resist /10 | Fair score | Model range | Status |",
        "|---:|---|---|---:|---:|---:|---:|---:|---:|---:|---|",
    ]
    for row in sorted(rows, key=lambda item: int(item["adjusted_rank"])):
        status = "Leader - moderate" if row["adjusted_rank"] == 1 else "Winner review" if row["winner_review"] else "Outside review group"
        report_lines.append(
            f"| {row['adjusted_rank']} | {row['code']} | {row['team']} | {row['current_total']:.2f} | {100 * row['judge'] * diagnostics['weights']['judge']:.2f} | {100 * row['adjusted_peer'] * diagnostics['weights']['peer']:.2f} | {100 * row['adjusted_attack'] * diagnostics['weights']['attack']:.2f} | {100 * row['adjusted_resistance'] * diagnostics['weights']['resistance']:.2f} | {row['adjusted_total']:.2f} | {row['adjusted_total_low']:.1f} - {row['adjusted_total_high']:.1f} | {status} |"
        )
    review_codes = ", ".join(
        str(row["code"])
        for row in sorted(rows, key=lambda item: int(item["adjusted_rank"]))
        if row["winner_review"]
    )
    report_lines.extend([
        "",
        "## Schedule and uncertainty",
        "",
        f"The estimated easiest schedule was {schedule_sorted[-1]['code']} ({schedule_sorted[-1]['schedule_easiness']:.1%} expected crack rate for an average attacker); the hardest was {schedule_sorted[0]['code']} ({schedule_sorted[0]['schedule_easiness']:.1%}).",
        "",
        "Largest ranking changes after adjustment:",
        "",
    ])
    for row in biggest_movers:
        report_lines.append(
            f"- {row['code']}: {row['current_rank']} to {row['adjusted_rank']} ({int(row['rank_change']):+d} places)."
        )
    report_lines.extend([
        "",
        f"Winner-review group: **{review_codes}**. Their modelled differences from the provisional leader include zero at the approximate 95% level.",
        "",
        "Intervals use conditional approximations at the selected regularisation strengths. They include model uncertainty in adjusted peer, cracking, and resistance scores; judge marks are treated as fixed. They do not account for model-selection uncertainty or the fact that the leader was selected from the same data.",
        "",
        "Because each team has few valid attempts, use the fair ordering as evidence and moderate the winner-review group with the published judging rubric. Do not use extra decimal places as a tie-break.",
        "",
        "## Competition rule for future rounds",
        "",
        "1. Record outcomes as `1 = cracked`, `0 = valid failure`, and blank or a separate status for broken/unsolvable attempts.",
        "2. Publish before competition that peer-rater and opponent-adjusted estimates will be used.",
        "3. Balance schedules using prior/pilot safe difficulty where possible; statistical adjustment is a fallback, not a substitute for good allocation.",
        "4. Require a minimum number of valid attempts or reduce the crack-component weight when data loss is high.",
        "5. Resolve near ties with a declared moderation rule, not extra decimal places from the model.",
    ])
    report_path.write_text("\n".join(report_lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    analysis_rows, analysis_attempts, analysis_diagnostics = analyse()
    write_outputs(analysis_rows, analysis_attempts, analysis_diagnostics)
    print(f"Wrote {OUTPUT.name} and {REPORT.name}")